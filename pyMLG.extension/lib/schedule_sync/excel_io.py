# -*- coding: utf-8 -*-
"""Excel-Ein- und Ausgabe über openpyxl.

Aufbau eines exportierten Blatts:
    Spalte A : UniqueId  (versteckt) - stabile Kennung für den Rückimport
    Spalte B : ElementId (nur zur Orientierung, beim Import ignoriert)
    Spalte C : Kategorie (nur zur Orientierung)
    Spalte D+: je ein Feld der Bauteilliste, Kopfzeile = Feldname

Zusätzlich wird ein verstecktes Blatt '_pyMLG_Meta' geschrieben, das je Spalte
die Parameterkennung (BuiltInParameter-Id bzw. Shared-Parameter-GUID) enthält.
Fehlt dieses Blatt beim Import, wird auf die Suche über den Feldnamen
zurückgefallen - der Import funktioniert also auch mit von Hand gebauten Dateien.
"""

import datetime
import json
import re

from schedule_sync.schedule_model import (
    ABSCHNITT_KOPF,
    ERSTE_DATENSPALTE,
    ERSTE_DATENZEILE,
    KOPFZEILE,
    SCHREIBBAR,
    SPALTE_ELEMENTID,
    SPALTE_KATEGORIE,
    SPALTE_UNIQUEID,
    STATUS_BERECHNET,
    STATUS_FEHLT,
    STATUS_GESPERRT,
    STATUS_TYP,
)

META_BLATT = u"_pyMLG_Meta"
META_VERSION = 1
KOPF_UNIQUEID = u"UniqueId"
KOPF_ELEMENTID = u"ElementId"
KOPF_KATEGORIE = u"Kategorie"
# Blätter mit der Tabelle wie in Revit - nur zum Lesen, der Import überspringt sie
ANSICHT_PRAEFIX = u"Ansicht - "

# In Excel-Blattnamen unzulässige Zeichen (Revit erlaubt sie in Ansichtsnamen)
UNGUELTIGE_BLATTZEICHEN = re.compile(r"[\[\]\:\*\?\/\\]")
# Steuerzeichen, an denen openpyxl mit IllegalCharacterError aussteigt
STEUERZEICHEN = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]")

MAX_SPALTENBREITE = 55
MIN_SPALTENBREITE = 10


# ---------------------------------------------------------------------------
# Hilfsfunktionen
# ---------------------------------------------------------------------------

def blattname(name, vergeben):
    """Revit-Ansichtsname -> gültiger, eindeutiger Excel-Blattname (max. 31 Zeichen)."""
    sauber = UNGUELTIGE_BLATTZEICHEN.sub(u"-", name or u"Tabelle").strip()
    sauber = sauber.strip(u"'") or u"Tabelle"
    sauber = sauber[:31]

    kandidat = sauber
    zaehler = 2
    while kandidat.lower() in vergeben:
        endung = u"_%d" % zaehler
        kandidat = sauber[:31 - len(endung)] + endung
        zaehler += 1
    vergeben.add(kandidat.lower())
    return kandidat


def _excel_wert(wert):
    """Macht einen beliebigen Wert für openpyxl schreibbar."""
    if wert is None:
        return None
    if isinstance(wert, bool):
        return u"Ja" if wert else u"Nein"
    if isinstance(wert, (int, float)):
        return wert
    text = wert if isinstance(wert, str) else str(wert)
    return STEUERZEICHEN.sub(u"", text)


def _breite(wert):
    """Geschätzte Zeichenbreite eines Zellwerts für die Spaltenbreite."""
    if wert is None:
        return 0
    return len(str(wert))


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

def schreibe_arbeitsmappe(openpyxl, pfad, bloecke, dokumentname,
                          blattschutz=True):
    """Schreibt alle Export-Blöcke in eine neue Arbeitsmappe.

    bloecke: Liste von Dicts mit den Schlüsseln
        name, uid, felder, zeilen, spaltenstatus
    Rückgabe: Liste (blattname, schedule_name, zeilenanzahl)
    """
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter

    kopf_font = Font(bold=True, color="FFFFFFFF")
    kopf_fuellung = PatternFill("solid", fgColor="FF44546A")
    kopf_gesperrt_fuellung = PatternFill("solid", fgColor="FF808080")
    id_fuellung = PatternFill("solid", fgColor="FFF2F2F2")
    grau = PatternFill("solid", fgColor="FFD9D9D9")
    orange = PatternFill("solid", fgColor="FFFCE4D6")
    gesperrt = Protection(locked=True)
    frei = Protection(locked=False)

    arbeitsmappe = openpyxl.Workbook()
    arbeitsmappe.remove(arbeitsmappe.active)

    meta = {
        "version": META_VERSION,
        "erzeugt": datetime.datetime.now().isoformat(timespec="seconds"),
        "dokument": dokumentname,
        "blaetter": {},
        "ansichten": [],
    }
    vergebene_namen = set()
    uebersicht = []

    for block in bloecke:
        name = blattname(block["name"], vergebene_namen)
        blatt = arbeitsmappe.create_sheet(title=name)
        felder = block["felder"]
        zeilen = block["zeilen"]
        status_je_spalte = block["spaltenstatus"]

        breiten = {}

        # --- Kopfzeile -----------------------------------------------------
        for spalte, titel in ((SPALTE_UNIQUEID, KOPF_UNIQUEID),
                              (SPALTE_ELEMENTID, KOPF_ELEMENTID),
                              (SPALTE_KATEGORIE, KOPF_KATEGORIE)):
            zelle = blatt.cell(row=KOPFZEILE, column=spalte, value=titel)
            zelle.font = kopf_font
            zelle.fill = kopf_fuellung
            zelle.protection = gesperrt
            breiten[spalte] = len(titel)

        for feld in felder:
            spalte = feld["spalte"]
            zelle = blatt.cell(row=KOPFZEILE, column=spalte, value=feld["name"])
            zelle.font = kopf_font
            zelle.alignment = Alignment(wrap_text=True, vertical="center")
            zelle.protection = gesperrt
            status = status_je_spalte.get(spalte, STATUS_GESPERRT)
            schreibbar = status in SCHREIBBAR
            zelle.fill = kopf_fuellung if schreibbar else kopf_gesperrt_fuellung

            notiz = _kopfnotiz(feld, status)
            if notiz:
                zelle.comment = Comment(notiz, u"pyMLG ScheduleSync")
            breiten[spalte] = len(feld["name"])

        # --- Datenzeilen ---------------------------------------------------
        for index, zeile in enumerate(zeilen):
            excel_zeile = ERSTE_DATENZEILE + index

            id_zellen = (
                (SPALTE_UNIQUEID, zeile["uid"]),
                (SPALTE_ELEMENTID, zeile["element_id"]),
                (SPALTE_KATEGORIE, zeile["kategorie"]),
            )
            for spalte, wert in id_zellen:
                zelle = blatt.cell(row=excel_zeile, column=spalte,
                                   value=_excel_wert(wert))
                zelle.fill = id_fuellung
                zelle.protection = gesperrt
                breiten[spalte] = max(breiten.get(spalte, 0), _breite(wert))

            for feld in felder:
                spalte = feld["spalte"]
                wert = _excel_wert(zeile["werte"].get(spalte))
                zelle = blatt.cell(row=excel_zeile, column=spalte, value=wert)
                status = zeile["status"].get(spalte, STATUS_GESPERRT)

                if status in SCHREIBBAR:
                    zelle.protection = frei
                    if status == STATUS_TYP:
                        # Typparameter wirken auf alle Instanzen dieses Typs
                        zelle.fill = orange
                else:
                    zelle.protection = gesperrt
                    zelle.fill = grau

                breiten[spalte] = max(breiten.get(spalte, 0), _breite(wert))

        # --- Darstellung ---------------------------------------------------
        blatt.freeze_panes = blatt.cell(row=ERSTE_DATENZEILE,
                                        column=ERSTE_DATENSPALTE).coordinate
        blatt.column_dimensions[get_column_letter(SPALTE_UNIQUEID)].hidden = True
        for spalte, breite in breiten.items():
            buchstabe = get_column_letter(spalte)
            blatt.column_dimensions[buchstabe].width = min(
                MAX_SPALTENBREITE, max(MIN_SPALTENBREITE, breite + 3))

        if zeilen:
            letzte_spalte = get_column_letter(
                max([SPALTE_KATEGORIE] + [f["spalte"] for f in felder]))
            blatt.auto_filter.ref = u"A%d:%s%d" % (
                KOPFZEILE, letzte_spalte, ERSTE_DATENZEILE + len(zeilen) - 1)

        if blattschutz:
            # Ohne Passwort: der Schutz ist eine Leitplanke, keine Sperre.
            # Gesperrte Zellen sind genau die, die beim Import ignoriert werden.
            blatt.protection.sheet = True
            blatt.protection.selectLockedCells = False
            blatt.protection.autoFilter = False
            blatt.protection.sort = False

        meta["blaetter"][name] = {
            "schedule_name": block["name"],
            "schedule_uid": block.get("uid"),
            "spalten": [{
                "spalte": feld["spalte"],
                "name": feld["name"],
                "feldname": feld.get("feldname") or feld["name"],
                "param_id": feld["param_id"],
                "guid": feld["guid"],
                "feldtyp": feld["feldtyp"],
                "ist_typfeld": feld["ist_typfeld"],
                "berechnet": feld["berechnet"],
                "einheit": feld["einheit"],
                "status": status_je_spalte.get(feld["spalte"], STATUS_GESPERRT),
            } for feld in felder],
        }
        uebersicht.append((name, block["name"], len(zeilen)))

        if block.get("ansicht"):
            ansicht_name = blattname(ANSICHT_PRAEFIX + block["name"],
                                     vergebene_namen)
            _schreibe_ansicht(arbeitsmappe, ansicht_name, block, blattschutz)
            meta["ansichten"].append(ansicht_name)

    meta_blatt = arbeitsmappe.create_sheet(title=META_BLATT)
    meta_blatt["A1"] = u"Technische Daten für den Rückimport - bitte nicht ändern."
    meta_blatt["A2"] = json.dumps(meta, ensure_ascii=False)
    meta_blatt.sheet_state = "hidden"

    arbeitsmappe.save(pfad)
    return uebersicht


def _schreibe_ansicht(arbeitsmappe, name, block, blattschutz):
    """Blatt mit der Tabelle genau so, wie Revit sie anzeigt (nur zum Lesen).

    Werte stehen als formatierter Text darin ('12,50 m²'), damit Gruppenköpfe,
    zusammengefasste Zeilen und Summen exakt der Revit-Darstellung entsprechen.
    """
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    blatt = arbeitsmappe.create_sheet(title=name)
    titel_font = Font(bold=True, size=12)
    kopf_font = Font(bold=True, color="FFFFFFFF")
    kopf_fuellung = PatternFill("solid", fgColor="FF44546A")
    spaltenkoepfe = set(feld["name"] for feld in block["felder"])
    breiten = {}

    for index, (abschnitt, texte) in enumerate(block["ansicht"]):
        zeile = index + 1
        belegte = [t for t in texte if t]
        # Spaltenkopfzeile erkennen: alle belegten Zellen sind Feldüberschriften
        ist_spaltenkopf = (abschnitt != ABSCHNITT_KOPF and len(belegte) > 1
                           and set(belegte) <= spaltenkoepfe)
        for spalte, text in enumerate(texte, start=1):
            zelle = blatt.cell(row=zeile, column=spalte,
                               value=_excel_wert(text) if text else None)
            if abschnitt == ABSCHNITT_KOPF:
                zelle.font = titel_font
            elif ist_spaltenkopf:
                zelle.font = kopf_font
                zelle.fill = kopf_fuellung
                zelle.alignment = Alignment(wrap_text=True, vertical="center")
            breiten[spalte] = max(breiten.get(spalte, 0), _breite(text))

    for spalte, breite in breiten.items():
        blatt.column_dimensions[get_column_letter(spalte)].width = min(
            MAX_SPALTENBREITE, max(MIN_SPALTENBREITE, breite + 3))

    blatt["A1"].comment = Comment(
        u"Tabelle wie in Revit - nur zur Ansicht.\n"
        u"Änderungen in diesem Blatt werden beim Import ignoriert. Bearbeiten "
        u"im Blatt '%s'." % block["name"], u"pyMLG ScheduleSync")
    blatt.sheet_properties.tabColor = "FF9AA5B1"
    if blattschutz:
        blatt.protection.sheet = True
        blatt.protection.selectLockedCells = False


def _kopfnotiz(feld, status):
    """Erklärender Kommentar an der Kopfzelle."""
    if status == STATUS_BERECHNET:
        return (u"Berechnetes Feld (%s).\n"
                u"Wird von Revit ermittelt und beim Import ignoriert."
                % (feld["feldtyp"] or u"?"))
    if status == STATUS_GESPERRT:
        return (u"Schreibgeschützter Parameter.\n"
                u"Änderungen in dieser Spalte werden beim Import ignoriert.")
    if status == STATUS_FEHLT:
        return u"Parameter an den Elementen nicht gefunden - wird ignoriert."
    if status == STATUS_TYP:
        return (u"Typparameter: Eine Änderung wirkt auf ALLE Instanzen "
                u"dieses Typs, nicht nur auf diese Zeile.")
    if feld.get("einheit"):
        return u"Einheit: %s (Projekteinheit)" % feld["einheit"]
    return None


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

def _lade_meta(arbeitsmappe):
    """Liest das versteckte Metablatt, falls vorhanden."""
    if META_BLATT not in arbeitsmappe.sheetnames:
        return None
    try:
        rohdaten = arbeitsmappe[META_BLATT]["A2"].value
        if not rohdaten:
            return None
        meta = json.loads(rohdaten)
        if meta.get("version") != META_VERSION:
            return None
        return meta
    except Exception:
        return None


def _felder_aus_meta(meta_blatt, kopfzeile):
    """Ordnet die Metadaten den tatsächlich vorhandenen Kopfzeilen zu.

    Die Zuordnung läuft über den Spaltennamen, nicht über die Spaltennummer -
    so bleiben eingefügte oder verschobene Spalten in Excel verkraftbar.
    """
    nach_name = {}
    if meta_blatt:
        for eintrag in meta_blatt.get("spalten", []):
            nach_name[eintrag.get("name")] = eintrag

    felder = []
    for spalte, titel in kopfzeile.items():
        if spalte < ERSTE_DATENSPALTE or not titel:
            continue
        eintrag = nach_name.get(titel)
        if eintrag is None:
            # Unbekannte Spalte: Suche später über den Feldnamen, Schreibschutz
            # wird zur Laufzeit am Parameter selbst geprüft.
            felder.append({
                "spalte": spalte,
                "name": titel,
                "feldname": titel,
                "param_id": None,
                "guid": None,
                "ist_typfeld": False,
                "berechnet": False,
                "status": None,
                "aus_meta": False,
            })
            continue
        felder.append({
            "spalte": spalte,
            "name": titel,
            "feldname": eintrag.get("feldname") or titel,
            "param_id": eintrag.get("param_id"),
            "guid": eintrag.get("guid"),
            "ist_typfeld": bool(eintrag.get("ist_typfeld")),
            "berechnet": bool(eintrag.get("berechnet")),
            "status": eintrag.get("status"),
            "aus_meta": True,
        })
    return felder


def lese_arbeitsmappe(openpyxl, pfad):
    """Liest eine exportierte (und bearbeitete) Arbeitsmappe.

    Rückgabe: (blaetter, warnungen)
    Blätter ohne 'UniqueId' in Spalte A werden mit Warnung übersprungen.
    """
    # data_only=True: Formeln werden als zuletzt berechneter Wert gelesen
    arbeitsmappe = openpyxl.load_workbook(pfad, data_only=True)
    meta = _lade_meta(arbeitsmappe)
    meta_blaetter = (meta or {}).get("blaetter", {})
    ansichten = set((meta or {}).get("ansichten", []))

    blaetter = []
    warnungen = []

    for name in arbeitsmappe.sheetnames:
        if name == META_BLATT:
            continue
        # Ansichtsblätter sind reine Lesekopien der Revit-Tabelle - still überspringen
        if name in ansichten or name.startswith(ANSICHT_PRAEFIX):
            continue
        blatt = arbeitsmappe[name]

        kopf_a = blatt.cell(row=KOPFZEILE, column=SPALTE_UNIQUEID).value
        if not kopf_a or str(kopf_a).strip() != KOPF_UNIQUEID:
            warnungen.append(
                u"Blatt '%s' übersprungen: Spalte A enthält keine Kopfzeile "
                u"'%s'." % (name, KOPF_UNIQUEID))
            continue

        kopfzeile = {}
        for spalte in range(1, blatt.max_column + 1):
            titel = blatt.cell(row=KOPFZEILE, column=spalte).value
            if titel is not None:
                kopfzeile[spalte] = str(titel).strip()

        felder = _felder_aus_meta(meta_blaetter.get(name), kopfzeile)
        if not felder:
            warnungen.append(
                u"Blatt '%s' übersprungen: keine Datenspalten gefunden." % name)
            continue

        zeilen = []
        for excel_zeile in range(ERSTE_DATENZEILE, blatt.max_row + 1):
            unique_id = blatt.cell(row=excel_zeile, column=SPALTE_UNIQUEID).value
            if unique_id is None or not str(unique_id).strip():
                continue  # Leerzeile am Ende oder vom Nutzer eingefügt
            werte = {}
            for feld in felder:
                werte[feld["spalte"]] = blatt.cell(
                    row=excel_zeile, column=feld["spalte"]).value
            zeilen.append({
                "excel_zeile": excel_zeile,
                "uid": str(unique_id).strip(),
                "werte": werte,
            })

        blaetter.append({
            "blattname": name,
            "schedule_name": (meta_blaetter.get(name) or {}).get(
                "schedule_name", name),
            "hat_meta": name in meta_blaetter,
            "felder": felder,
            "zeilen": zeilen,
        })

    if meta is None:
        warnungen.append(
            u"Kein Metablatt gefunden - die Parameter werden anhand der "
            u"Spaltennamen gesucht. Das funktioniert, ist aber weniger eindeutig.")

    return blaetter, warnungen
